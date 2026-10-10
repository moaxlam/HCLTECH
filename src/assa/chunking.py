"""Document reading and clause-aware chunking.

Chunking by *clause* rather than by a fixed character window is what makes R2
(citations on every factual answer) cheap: every chunk knows the document, the
clause number, the page and the version, so a citation can be built from it
without guessing.

Conventions understood by the chunker:

* ``<!-- page: 12 -->`` marks the page that the text after it appears on.
* Markdown headings give the heading path (``## 7. Attendance and Detention``).
* A paragraph that starts with a clause number (``7.2 Minimum attendance ...``)
  starts a new chunk with ``section = "7.2"``.
* A paragraph that starts with ``**Q7. ...**`` starts a new chunk with
  ``section = "Q7"`` (help-desk FAQ style).

Long clauses are split on paragraph and sentence boundaries with a small overlap
so no fact is orphaned at a chunk edge.
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

PAGE_MARKER_RE = re.compile(r"<!--\s*page:\s*(\d+)\s*-->", re.IGNORECASE)
CLAUSE_RE = re.compile(r"^(\d+(?:\.\d+)*)[.)]?\s+\S")
FAQ_RE = re.compile(r"^\*\*(Q\d+)[.)]\s*(.+?)\*\*", re.IGNORECASE | re.DOTALL)
HEADING_RE = re.compile(r"^(#{1,6})\s+(.*)$", re.MULTILINE)

SUPPORTED_SUFFIXES = {".md", ".markdown", ".txt", ".text", ".pdf", ".docx"}


def ordinal_date(value: Any) -> int:
    """``YYYY-MM-DD`` -> ``YYYYMMDD``; ``0`` when there is no usable date."""
    text = str(value or "").strip()[:10]
    digits = text.replace("-", "")
    return int(digits) if len(digits) == 8 and digits.isdigit() else 0


@dataclass
class Chunk:
    chunk_id: str
    doc_id: str
    text: str
    section: str
    heading: str
    page: int | None
    doc_metadata: dict[str, Any] = field(default_factory=dict)

    def as_record(self) -> tuple[str, str, dict[str, Any]]:
        """``(chunk_id, text, metadata)`` for the vector store."""
        metadata = {
            "doc_id": self.doc_id,
            "section": self.section,
            "heading": self.heading,
            "page": self.page if self.page is not None else -1,
            "title": str(self.doc_metadata.get("title", "")),
            "version": str(self.doc_metadata.get("version", "")),
            "effective_from": str(self.doc_metadata.get("effective_from", "")),
            "effective_to": str(self.doc_metadata.get("effective_to", "")),
            # Numeric sort key for the same date. The vector store rejects string
            # operands for range operators, and R4 needs a "takes effect later"
            # range filter, so the date is indexed twice.
            "effective_from_ord": ordinal_date(self.doc_metadata.get("effective_from")),
            "authority_level": int(self.doc_metadata.get("authority_level", 3) or 3),
            "doc_type": str(self.doc_metadata.get("doc_type", "")),
            "issuer": str(self.doc_metadata.get("issuer", "")),
            "supersedes": str(self.doc_metadata.get("supersedes", "")),
            "scope_programmes": str(self.doc_metadata.get("scope_programmes", "ALL")),
            "scope_batches": str(self.doc_metadata.get("scope_batches", "ALL")),
            "synthetic": str(self.doc_metadata.get("synthetic", "N")),
        }
        for key, value in metadata.items():
            if value == "":
                metadata[key] = "-"
        return self.chunk_id, self.text, metadata


@dataclass
class ReadResult:
    text: str
    needs_ocr: bool = False
    note: str = ""


def read_document(path: Path) -> ReadResult:
    """Read a document to text. PDFs use pypdf; scanned pages report ``needs_ocr``."""
    suffix = path.suffix.lower()
    if suffix in {".md", ".markdown", ".txt", ".text"}:
        return ReadResult(text=path.read_text(encoding="utf-8", errors="replace"))

    if suffix == ".docx":
        return ReadResult(text=_read_docx(path))

    if suffix == ".pdf":
        text, pages = _read_pdf(path)
        if len(text.strip()) < 200 and pages:
            ocr_text = _try_ocr(path)
            if ocr_text:
                return ReadResult(text=ocr_text, note="text recovered with OCR")
            return ReadResult(
                text=text,
                needs_ocr=True,
                note="the PDF contains little extractable text; install the OCR extras (see README)",
            )
        return ReadResult(text=text)

    raise ValueError(f"unsupported document type {suffix!r} (supported: {sorted(SUPPORTED_SUFFIXES)})")


def _read_docx(path: Path) -> str:
    try:
        import docx  # python-docx
    except ImportError as exc:  # pragma: no cover - optional dependency
        raise ValueError("python-docx is required to read .docx files") from exc
    document = docx.Document(str(path))
    lines = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            lines.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n\n".join(line for line in lines if line is not None)


def _read_pdf(path: Path) -> tuple[str, int]:
    try:
        from pypdf import PdfReader
    except ImportError:
        try:
            from PyPDF2 import PdfReader  # type: ignore
        except ImportError as exc:
            raise ValueError("pypdf is required to read .pdf files") from exc
    reader = PdfReader(str(path))
    pieces = []
    for index, page in enumerate(reader.pages, start=1):
        pieces.append(f"<!-- page: {index} -->\n{page.extract_text() or ''}")
    return "\n\n".join(pieces), len(reader.pages)


def _try_ocr(path: Path) -> str:
    """Optional OCR path for scanned pages (pytesseract + pdf2image)."""
    try:
        import pytesseract  # type: ignore
        from pdf2image import convert_from_path  # type: ignore
    except ImportError:
        return ""
    try:
        images = convert_from_path(str(path))
    except Exception:
        return ""
    pieces = []
    for index, image in enumerate(images, start=1):
        pieces.append(f"<!-- page: {index} -->\n{pytesseract.image_to_string(image)}")
    return "\n\n".join(pieces)


def _split_paragraphs(text: str) -> list[tuple[str, int | None, str]]:
    """Return ``(paragraph, page, heading_path)`` triples."""
    paragraphs: list[tuple[str, int | None, str]] = []
    page: int | None = None
    heading_path = ""
    buffer: list[str] = []

    def flush() -> None:
        if not buffer:
            return
        block = "\n".join(buffer).strip()
        buffer.clear()
        if block:
            paragraphs.append((block, page, heading_path))

    for raw_line in text.splitlines():
        marker = PAGE_MARKER_RE.search(raw_line)
        if marker:
            flush()
            page = int(marker.group(1))
            remainder = PAGE_MARKER_RE.sub("", raw_line).strip()
            if remainder:
                buffer.append(remainder)
            continue
        heading = HEADING_RE.match(raw_line.rstrip())
        if heading:
            flush()
            level = len(heading.group(1))
            title = heading.group(2).strip()
            parts = heading_path.split(" > ") if heading_path else []
            parts = parts[: level - 1]
            heading_path = " > ".join(parts + [title]) if title else heading_path
            buffer.append(raw_line.rstrip())
            continue
        if raw_line.strip() == "":
            flush()
            continue
        buffer.append(raw_line.rstrip())
    flush()
    return paragraphs


def _clause_of(paragraph: str) -> str | None:
    first_line = paragraph.lstrip().splitlines()[0] if paragraph.strip() else ""
    faq = FAQ_RE.match(first_line)
    if faq:
        return faq.group(1).upper()
    clause = CLAUSE_RE.match(first_line)
    if clause:
        return clause.group(1)
    return None


def chunk_document(
    text: str,
    doc_metadata: dict[str, Any],
    max_chars: int = 1400,
    overlap_chars: int = 150,
    min_chars: int = 120,
) -> list[Chunk]:
    doc_id = str(doc_metadata.get("doc_id", "UNKNOWN"))
    paragraphs = _split_paragraphs(text)

    blocks: list[tuple[str, str, str, int | None]] = []  # (section, heading, text, page)
    current_section = ""
    current_heading = ""
    current_page: int | None = None
    buffer: list[str] = []

    def flush() -> None:
        if buffer:
            blocks.append((current_section, current_heading, "\n".join(buffer).strip(), current_page))
            buffer.clear()

    for paragraph, page, heading_path in paragraphs:
        clause = _clause_of(paragraph)
        is_heading_only = paragraph.lstrip().startswith("#")
        if clause and not is_heading_only:
            flush()
            current_section = clause
        if is_heading_only and not current_section:
            current_heading = heading_path
        elif heading_path:
            current_heading = heading_path
        if not buffer:
            current_page = page
        buffer.append(paragraph)
        flush_if_large = sum(len(part) for part in buffer)
        if flush_if_large > max_chars:
            flush()
    flush()

    merged: list[tuple[str, str, str, int | None]] = []
    for section, heading, block_text, page in blocks:
        if merged and len(merged[-1][2]) < min_chars and section == merged[-1][0]:
            previous = merged.pop()
            merged.append((previous[0], previous[1], previous[2] + "\n\n" + block_text, previous[3]))
        else:
            merged.append((section, heading, block_text, page))

    chunks: list[Chunk] = []
    for index, (section, heading, block_text, page) in enumerate(merged):
        label = section or (heading.split(" > ")[-1] if heading else f"part-{index + 1}")
        for piece_index, piece in enumerate(_split_long(block_text, max_chars, overlap_chars)):
            suffix = f"#{piece_index + 1}" if piece_index else ""
            chunks.append(
                Chunk(
                    chunk_id=f"{doc_id}::{label}::{piece_index}",
                    doc_id=doc_id,
                    text=piece,
                    section=label,
                    heading=heading,
                    page=page,
                    doc_metadata=doc_metadata,
                )
            )
    return [chunk for chunk in chunks if chunk.text.strip() and not _is_heading_only(chunk.text)]


def _is_heading_only(text: str) -> bool:
    """A block that is nothing but markdown headings carries no answerable content."""
    lines = [line.strip() for line in text.splitlines() if line.strip()]
    return bool(lines) and all(line.startswith("#") for line in lines)


def _split_long(text: str, max_chars: int, overlap_chars: int) -> list[str]:
    if len(text) <= max_chars:
        return [text]
    pieces: list[str] = []
    paragraphs = [part for part in text.split("\n\n") if part.strip()]
    current = ""
    for paragraph in paragraphs:
        if len(current) + len(paragraph) + 2 <= max_chars:
            current = f"{current}\n\n{paragraph}".strip()
            continue
        if current:
            pieces.append(current)
        if len(paragraph) <= max_chars:
            current = paragraph
            continue
        sentences = re.split(r"(?<=[.;])\s+", paragraph)
        current = ""
        for sentence in sentences:
            if len(current) + len(sentence) + 1 <= max_chars:
                current = f"{current} {sentence}".strip()
            else:
                if current:
                    pieces.append(current)
                current = sentence
        if current:
            pieces.append(current)
        current = ""
    if current:
        pieces.append(current)
    if overlap_chars > 0 and len(pieces) > 1:
        overlapped = [pieces[0]]
        for index in range(1, len(pieces)):
            tail = pieces[index - 1][-overlap_chars:]
            overlapped.append(f"{tail}\n{pieces[index]}".strip())
        pieces = overlapped
    return pieces
